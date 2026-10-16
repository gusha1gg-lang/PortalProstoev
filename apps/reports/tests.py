"""Тесты дашбордов и отчётов (запуск: python manage.py test apps.reports)."""
import io
import json
from datetime import datetime
from decimal import Decimal

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone

from apps.catalog.models import Enterprise, ItFaultCategory, LossRate, Section, Workshop
from apps.downtimes.models import DowntimeEvent
from apps.downtimes.services import create_case_from_events

from .services import incident_rows, kpi, monthly_series, top_categories


def dt(text):
    """Дата-время с часовым поясом из строки «ГГГГ-ММ-ДД ЧЧ:ММ»."""
    return timezone.make_aware(datetime.strptime(text, "%Y-%m-%d %H:%M"))


class ReportsTestBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.enterprise = Enterprise.objects.create(name='Тестовый завод')
        cls.workshop = Workshop.objects.create(enterprise=cls.enterprise, name="ТЭСЦ-4")
        cls.section = Section.objects.create(workshop=cls.workshop, name="Участок отделки")
        LossRate.objects.create(
            section=cls.section,
            cost_per_hour=Decimal("120000"),
            valid_from="2025-01-01",
        )
        cls.category = ItFaultCategory.objects.create(code="net", name="Сеть")
        cls.user = get_user_model().objects.create_user(username="tester", password="TestPass123!")

    def make_event(self, start, end, side=DowntimeEvent.FaultSide.IT, category=None):
        return DowntimeEvent.objects.create(
            started_at=dt(start),
            ended_at=dt(end),
            workshop=self.workshop,
            section=self.section,
            fault_side=side,
            fault_category=category,
            status=DowntimeEvent.Status.CONFIRMED,
        )


class KpiTests(ReportsTestBase):
    def test_kpi_totals_and_loss(self):
        self.make_event("2026-10-10 15:00", "2026-10-10 15:30", category=self.category)
        self.make_event("2026-10-11 10:00", "2026-10-11 10:15", category=self.category)
        data = kpi(2026, 10)

        self.assertEqual(data["incident_count"], 2)
        self.assertEqual(data["downtime_min"], 45)
        # 0,5 ч × 120 000 + 0,25 ч × 120 000 = 60 000 + 30 000
        self.assertEqual(data["loss"], Decimal("90000.00"))

    def test_it_share_accounts_for_other_services(self):
        # 30 минут вины ИТ + 30 минут другой службы → доля ИТ 50%.
        self.make_event("2026-10-10 15:00", "2026-10-10 15:30", side=DowntimeEvent.FaultSide.IT)
        self.make_event("2026-10-11 15:00", "2026-10-11 15:30", side=DowntimeEvent.FaultSide.FOREIGN)
        data = kpi(2026, 10)
        self.assertEqual(data["downtime_min"], 30)  # потери считаем только по ИТ
        self.assertEqual(data["it_share"], Decimal("50.00"))

    def test_incidents_carry_category_and_enterprise(self):
        self.make_event("2026-10-10 15:00", "2026-10-10 15:30", category=self.category)
        data = kpi(2026, 10)
        incident = data["incidents"][0]
        self.assertEqual(incident["category"], "Сеть")
        self.assertEqual(incident["enterprise"], "Тестовый завод")
        self.assertEqual(incident["kind"], "event")

    def test_merged_case_counts_once(self):
        """Дубли, собранные в кейс, на дашборде считаются одним инцидентом."""
        e1 = self.make_event("2026-10-10 15:05", "2026-10-10 15:20")
        e2 = self.make_event("2026-10-10 15:05", "2026-10-10 15:20")
        create_case_from_events([e1, e2], user=None)

        data = kpi(2026, 10)
        # Два события по 15 минут внахлёст → один инцидент на 15 минут.
        self.assertEqual(data["incident_count"], 1)
        self.assertEqual(data["downtime_min"], 15)
        self.assertEqual(data["incidents"][0]["kind"], "case")


class SliceTests(ReportsTestBase):
    def test_top_categories_sorted_by_minutes(self):
        self.make_event("2026-10-10 15:00", "2026-10-10 15:30", category=self.category)
        self.make_event("2026-10-11 15:00", "2026-10-11 15:10")  # без категории
        data = kpi(2026, 10)
        categories = top_categories(data["incidents"])
        self.assertEqual(categories[0]["name"], "Сеть")
        self.assertEqual(categories[0]["minutes"], 30)
        self.assertIn("Без категории", [item["name"] for item in categories])

    def test_monthly_series_has_requested_length(self):
        series = monthly_series(2026, 3, months=6)
        self.assertEqual(len(series), 6)
        # Идём от старого к новому: последний месяц — запрошенный.
        self.assertEqual(series[-1]["label"], "03.2026")
        self.assertEqual(series[0]["label"], "10.2025")

    def test_incident_rows_for_export(self):
        self.make_event("2026-10-10 15:00", "2026-10-10 15:30", category=self.category)
        rows = incident_rows(2026, 10)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["period"], "2026-10")
        self.assertEqual(rows[0]["category"], "Сеть")


class ViewTests(ReportsTestBase):
    def setUp(self):
        # Страницы закрыты входом (этап 9) — входим тестовым пользователем.
        self.client.force_login(self.user)

    def test_anonymous_is_redirected_to_login(self):
        self.client.logout()
        response = self.client.get("/")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/accounts/login/", response["Location"])

    def test_dashboard_renders(self):
        self.make_event("2026-10-10 15:00", "2026-10-10 15:30", category=self.category)
        response = self.client.get("/", {"period": "2026-10"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Дашборд")
        self.assertContains(response, "Сеть")

    def test_sla_report_renders(self):
        response = self.client.get("/reports/sla/", {"period": "2026-10"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "SLA-отчёт")

    def test_export_csv_has_bom_and_header(self):
        self.make_event("2026-10-10 15:00", "2026-10-10 15:30")
        response = self.client.get("/reports/export/", {"period": "2026-10"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])
        body = response.content.decode("utf-8-sig")
        self.assertIn("Предприятие", body)
        self.assertIn("Тестовый завод", body)

    def test_export_xlsx_is_valid(self):
        self.make_event("2026-10-10 15:00", "2026-10-10 15:30")
        response = self.client.get("/reports/export/xlsx/", {"period": "2026-10"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet.cell(row=1, column=1).value, "Период")
        self.assertEqual(sheet.max_row, 2)  # шапка + 1 инцидент

    def test_chart_data_is_json(self):
        self.make_event("2026-10-10 15:00", "2026-10-10 15:30")
        response = self.client.get("/reports/chart-data/", {"period": "2026-10"})
        self.assertEqual(response.status_code, 200)
        payload = json.loads(response.content)
        self.assertIn("series", payload)
        self.assertEqual(len(payload["series"]), 6)
